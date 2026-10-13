from pptx import Presentation

prs = Presentation("novaretail_template.pptx")

for slide_no, slide in enumerate(prs.slides, start=1):

    print(f"\nSLIDE {slide_no}")

    for shape in slide.shapes:
        print(shape.name)